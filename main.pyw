from PyQt5 import QtWidgets, uic
import sys
import docx
import subprocess
import generators

class Ui(QtWidgets.QMainWindow):
    def __init__(self):
        super(Ui, self).__init__()
        uic.loadUi('bars.ui', self)
        self.ROW = 0
        self.invoices=get_from_db()
        for invoice in reversed(self.invoices):
            #print(invoice)
            self.invoiceBox.addItem(invoice)

        self.invoiceBox.activated.connect(self.set_customer_name)
        self.searchBtn.clicked.connect(self.hosps_list)
        self.taskBtn.clicked.connect(self.display_task)
        self.OkBtn.clicked.connect(self.BarcodeGeneration)
        self.show()

    def set_customer_name(self):
        if get_from_db((self.invoiceBox.currentText())[0] is None:
            self.customerEdit.setText('Отсутств.')
        else:
            st = get_from_db((self.invoiceBox.currentText())[0]
            i = st.find(',')
            if i != -1:
                stri = st[:i]
            self.customerEdit.setText(stri)

    def BarcodeGeneration(self):
        result = UritGen('ООО "БелАнта"', '2337069235', 'LICG900V013021A-81697')
        if result != 'UritGen - Нечего выполнять':
            os.system(str(source_dir)+'\Bases\out.dbf')
        else: print(result)
        #self.close()
        #source_dir = Path.cwd()
        
#        HOSP =table.cell(self.ROW,0).text
#        UID = table.cell(self.ROW,2).text
#        SN = table.cell(self.ROW,3).text
#        print(HOSP,UID,SN)
#        if self.WhichGenerator == 0:
#            strcode = TecomGen(HOSP,UID,SN)
#            print(strcode)
#            if strcode != 'TecomRobot - Нечего выполнять!':
#                   os.system(str(source_dir)+'\Bases\\outTecom.dbf')
###                    subprocess.call('C:\\Users\\AM\\Desktop\\LVWIN60\\Lvdbed.exe C:\\Users\\AM\\Desktop\\Bars\\outblab.dbf')
###                    subprocess.call('C:\\Users\\AM\\Desktop\\LVWIN60\\lv.exe')      
#        elif self.WhichGenerator == 3:
#            strcode = UritGen(HOSP,UID,SN)
#            print(strcode)
#            if strcode != 'UritGen - Нечего выполнять!':
#                   os.system(str(source_dir)+'\Bases\out.dbf')
#        else:
#            print('Непонятно что делать')
            
        
    
    def hosps_list(self):
        self.HospsTable.setColumnCount(3)                               # Устанавливаем кол-во колонок
        self.HospsTable.setRowCount(len(newtable))                      # и строк в таблице
        self.HospsTable.setHorizontalHeaderLabels(['Тип прибора','Клиника','SN']) # заголовок таблицы
                for r in range(len(newtable)):
            for c in range(3):
                #print(r,c-1)
                item = QtWidgets.QTableWidgetItem(newtable[r][c], 0)
                self.HospsTable.setItem(r, c, item)
        self.HospsTable.resizeColumnsToContents()
        #self.HospsTable.cellClicked.connect(self.on_click)


    def on_click(self):
        self.ROW = self.HospsTable.currentRow()+1


    def display_task(self):
        pass


app = QtWidgets.QApplication(sys.argv)
window = Ui()
app.exec_()
